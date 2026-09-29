import os, secrets, io, smtplib
from email.message import EmailMessage
from datetime import datetime, date, timedelta
from flask import Flask, jsonify, request, send_file
from flask_cors import CORS
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import text
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import load_workbook, Workbook

db=SQLAlchemy()

class Member(db.Model):
    __tablename__="members"
    id=db.Column(db.Integer,primary_key=True)
    first_name=db.Column(db.String(100),nullable=False)
    last_name=db.Column(db.String(100),nullable=False)
    company=db.Column(db.String(180))
    activity=db.Column(db.String(180))
    email=db.Column(db.String(180),index=True)
    phone=db.Column(db.String(50))
    created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)

class UserAccount(db.Model):
    __tablename__="user_accounts"
    id=db.Column(db.Integer,primary_key=True)
    member_id=db.Column(db.Integer,db.ForeignKey("members.id"),nullable=False,unique=True,index=True)
    email=db.Column(db.String(180),nullable=False,unique=True,index=True)
    password_hash=db.Column(db.String(255))
    status=db.Column(db.String(30),nullable=False,default="INVITED")
    invitation_token=db.Column(db.String(128),unique=True,index=True)
    invitation_expires_at=db.Column(db.DateTime)
    activated_at=db.Column(db.DateTime)
    is_system_admin=db.Column(db.Boolean,nullable=False,default=False)
    member=db.relationship("Member",backref=db.backref("user_account",uselist=False))

class AuthSession(db.Model):
    __tablename__="auth_sessions"
    id=db.Column(db.Integer,primary_key=True)
    account_id=db.Column(db.Integer,db.ForeignKey("user_accounts.id"),nullable=False,index=True)
    token_hash=db.Column(db.String(64),nullable=False,unique=True,index=True)
    expires_at=db.Column(db.DateTime,nullable=False)
    created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)
    account=db.relationship("UserAccount")

class Membership(db.Model):
    __tablename__="memberships"
    id=db.Column(db.Integer,primary_key=True)
    member_id=db.Column(db.Integer,db.ForeignKey("members.id"),nullable=False,index=True)
    start_date=db.Column(db.Date)
    end_date=db.Column(db.Date)
    status=db.Column(db.String(30),nullable=False,default="ACTIVE")
    member=db.relationship("Member",backref="memberships")

class AuditLog(db.Model):
    __tablename__="audit_logs"
    id=db.Column(db.Integer,primary_key=True)
    account_id=db.Column(db.Integer,db.ForeignKey("user_accounts.id"),index=True)
    member_id=db.Column(db.Integer,db.ForeignKey("members.id"),index=True)
    action=db.Column(db.String(100),nullable=False,index=True)
    entity_type=db.Column(db.String(80),index=True)
    entity_id=db.Column(db.String(80))
    details=db.Column(db.Text)
    created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False,index=True)

class HandoverItem(db.Model):
    __tablename__="handover_items"
    id=db.Column(db.Integer,primary_key=True)
    from_mandate_id=db.Column(db.Integer,db.ForeignKey("mandates.id"),nullable=False,index=True)
    to_mandate_id=db.Column(db.Integer,db.ForeignKey("mandates.id"),nullable=False,index=True)
    role_id=db.Column(db.Integer,db.ForeignKey("roles.id"),nullable=False,index=True)
    title=db.Column(db.String(255),nullable=False)
    notes=db.Column(db.Text)
    outgoing_confirmed_at=db.Column(db.DateTime)
    incoming_confirmed_at=db.Column(db.DateTime)
    created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)

class Resource(db.Model):
    __tablename__="resources"
    id=db.Column(db.Integer,primary_key=True)
    title=db.Column(db.String(255),nullable=False)
    description=db.Column(db.Text)
    category=db.Column(db.String(60),nullable=False,default="GENERAL")
    url=db.Column(db.Text)
    role_id=db.Column(db.Integer,db.ForeignKey("roles.id"),nullable=True,index=True)
    board_only=db.Column(db.Boolean,nullable=False,default=False)
    active=db.Column(db.Boolean,nullable=False,default=True)
    file_name=db.Column(db.String(255))
    file_type=db.Column(db.String(120))
    file_data=db.Column(db.LargeBinary)
    created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)

class Platform(db.Model):
    __tablename__="platforms"
    id=db.Column(db.Integer,primary_key=True)
    name=db.Column(db.String(180),nullable=False)
    description=db.Column(db.Text)
    url=db.Column(db.Text)
    owner_role_id=db.Column(db.Integer,db.ForeignKey("roles.id"),nullable=True,index=True)
    board_only=db.Column(db.Boolean,nullable=False,default=False)
    renewal_info=db.Column(db.String(255))
    active=db.Column(db.Boolean,nullable=False,default=True)

class Mentorship(db.Model):
    __tablename__="mentorships"
    id=db.Column(db.Integer,primary_key=True)
    member_id=db.Column(db.Integer,db.ForeignKey("members.id"),nullable=False,index=True)
    mentor_id=db.Column(db.Integer,db.ForeignKey("members.id"),nullable=False,index=True)
    status=db.Column(db.String(30),nullable=False,default="ACTIVE")
    started_at=db.Column(db.Date,default=date.today)
    completed_at=db.Column(db.Date)

class OnboardingStep(db.Model):
    __tablename__="onboarding_steps"
    id=db.Column(db.Integer,primary_key=True)
    member_id=db.Column(db.Integer,db.ForeignKey("members.id"),nullable=False,index=True)
    code=db.Column(db.String(60),nullable=False)
    label=db.Column(db.String(180),nullable=False)
    completed=db.Column(db.Boolean,nullable=False,default=False)
    completed_at=db.Column(db.DateTime)

class Mandate(db.Model):
    __tablename__="mandates"
    id=db.Column(db.Integer,primary_key=True)
    name=db.Column(db.String(120),nullable=False)
    start_date=db.Column(db.Date,nullable=False)
    end_date=db.Column(db.Date,nullable=False)
    status=db.Column(db.String(30),nullable=False,default="PREPARATION")

class Role(db.Model):
    __tablename__="roles"
    id=db.Column(db.Integer,primary_key=True)
    name=db.Column(db.String(120),nullable=False,unique=True)

class RoleAssignment(db.Model):
    __tablename__="role_assignments"
    id=db.Column(db.Integer,primary_key=True)
    mandate_id=db.Column(db.Integer,db.ForeignKey("mandates.id"),nullable=False)
    role_id=db.Column(db.Integer,db.ForeignKey("roles.id"),nullable=False)
    member_id=db.Column(db.Integer,db.ForeignKey("members.id"),nullable=False)
    status=db.Column(db.String(30),nullable=False,default="ACTIVE")

class ImportLog(db.Model):
    __tablename__="import_logs"
    id=db.Column(db.Integer,primary_key=True)
    account_id=db.Column(db.Integer,db.ForeignKey("user_accounts.id"),nullable=False,index=True)
    filename=db.Column(db.String(255))
    created_count=db.Column(db.Integer,nullable=False,default=0)
    updated_count=db.Column(db.Integer,nullable=False,default=0)
    error_count=db.Column(db.Integer,nullable=False,default=0)
    created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)
    account=db.relationship("UserAccount")

class Duty(db.Model):
    __tablename__="duties"
    id=db.Column(db.Integer,primary_key=True)
    code=db.Column(db.String(40),nullable=False,unique=True)
    name=db.Column(db.String(120),nullable=False)
    active=db.Column(db.Boolean,nullable=False,default=True)

class DutyAssignment(db.Model):
    __tablename__="duty_assignments"
    id=db.Column(db.Integer,primary_key=True)
    meeting_id=db.Column(db.Integer,db.ForeignKey("meetings.id"),nullable=False,index=True)
    duty_id=db.Column(db.Integer,db.ForeignKey("duties.id"),nullable=False,index=True)
    member_id=db.Column(db.Integer,db.ForeignKey("members.id"),nullable=False,index=True)
    status=db.Column(db.String(30),nullable=False,default="DRAFT")
    created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)

class Notification(db.Model):
    __tablename__="notifications"
    id=db.Column(db.Integer,primary_key=True)
    member_id=db.Column(db.Integer,db.ForeignKey("members.id"),nullable=False,index=True)
    type=db.Column(db.String(50),nullable=False)
    title=db.Column(db.String(255),nullable=False)
    message=db.Column(db.Text)
    link=db.Column(db.String(120))
    read_at=db.Column(db.DateTime)
    created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False,index=True)

class RequestRecipient(db.Model):
    __tablename__="request_recipients"
    id=db.Column(db.Integer,primary_key=True)
    request_type=db.Column(db.String(40),nullable=False,index=True)
    member_id=db.Column(db.Integer,db.ForeignKey("members.id"),nullable=False,index=True)
    created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)

class MemberRequest(db.Model):
    __tablename__="member_requests"
    id=db.Column(db.Integer,primary_key=True)
    member_id=db.Column(db.Integer,db.ForeignKey("members.id"),nullable=False,index=True)
    type=db.Column(db.String(40),nullable=False,index=True)
    title=db.Column(db.String(255))
    message=db.Column(db.Text)
    guest_name=db.Column(db.String(255))
    guest_email=db.Column(db.String(255))
    event_date=db.Column(db.Date)
    attachment_name=db.Column(db.String(255))
    attachment_type=db.Column(db.String(120))
    attachment_data=db.Column(db.LargeBinary)
    status=db.Column(db.String(30),nullable=False,default="TO_PROCESS")
    created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)

class SwapRequest(db.Model):
    __tablename__="swap_requests"
    id=db.Column(db.Integer,primary_key=True)
    assignment_id=db.Column(db.Integer,db.ForeignKey("duty_assignments.id"),nullable=False,index=True)
    requester_id=db.Column(db.Integer,db.ForeignKey("members.id"),nullable=False,index=True)
    volunteer_id=db.Column(db.Integer,db.ForeignKey("members.id"),index=True)
    status=db.Column(db.String(30),nullable=False,default="OPEN")
    created_at=db.Column(db.DateTime,default=datetime.utcnow,nullable=False)

class Meeting(db.Model):
    __tablename__="meetings"
    id=db.Column(db.Integer,primary_key=True)
    date=db.Column(db.Date,nullable=False,index=True)
    start_time=db.Column(db.Time)
    type=db.Column(db.String(40),nullable=False,default="REGULAR")
    location=db.Column(db.String(180))
    status=db.Column(db.String(30),nullable=False,default="PLANNED")
    notes=db.Column(db.Text)

def create_app():
    app=Flask(__name__)
    database_url=os.getenv("DATABASE_URL","sqlite:///bni_dev.db")
    if database_url.startswith("postgres://"):
        database_url=database_url.replace("postgres://","postgresql://",1)
    app.config["SQLALCHEMY_DATABASE_URI"]=database_url
    app.config["SQLALCHEMY_TRACK_MODIFICATIONS"]=False
    db.init_app(app)
    CORS(app,origins=os.getenv("CORS_ORIGINS","https://cx-labs.be").split(","))

    @app.get("/health")
    def health():
        return {"status":"ok","service":"bni-soignies-api"}

    @app.get("/api")
    def api():
        return {"name":"BNI Soignies Hub API","version":"0.1.0"}

    def audit(account,action,entity_type=None,entity_id=None,details=None):
        try:
            db.session.add(AuditLog(account_id=account.id if account else None,member_id=account.member_id if account else None,action=action,entity_type=entity_type,entity_id=str(entity_id) if entity_id is not None else None,details=details))
        except Exception:
            pass

    def active_mandate(on_date=None):
        d=on_date or date.today()
        return Mandate.query.filter(Mandate.start_date<=d,Mandate.end_date>=d).order_by(Mandate.start_date.desc()).first()

    def current_assignment(member_id,role_id=None,on_date=None):
        mandate=active_mandate(on_date)
        if not mandate:return None
        q=RoleAssignment.query.filter_by(member_id=member_id,mandate_id=mandate.id,status="ACTIVE")
        if role_id is not None:q=q.filter_by(role_id=role_id)
        return q.first()

    def require_admin():
        account,_=bearer_account()
        if not account:return None,({"error":"Authentification requise."},401)
        if not account_payload(account)["isAdmin"]:return None,({"error":"Accès administrateur requis."},403)
        return account,None

    def member_payload(m):
        membership=Membership.query.filter_by(member_id=m.id).order_by(Membership.id.desc()).first()
        account=m.user_account
        return {"id":m.id,"firstName":m.first_name,"lastName":m.last_name,"company":m.company,"activity":m.activity,"email":m.email,"phone":m.phone,"status":membership.status if membership else None,"startDate":membership.start_date.isoformat() if membership and membership.start_date else None,"endDate":membership.end_date.isoformat() if membership and membership.end_date else None,"accountStatus":account.status if account else "NO_ACCOUNT","isSystemAdmin":bool(account and account.is_system_admin)}

    @app.route("/api/members",methods=["GET","POST"])
    def members():
        _,auth_error=require_admin()
        if auth_error:return auth_error
        if request.method=="POST":
            data=request.get_json(silent=True) or {}
            first=(data.get("firstName") or "").strip()
            last=(data.get("lastName") or "").strip()
            if not first or not last:
                return {"error":"firstName and lastName are required"},400
            email=(data.get("email") or "").strip() or None
            duplicate=Member.query.filter(db.func.lower(Member.first_name)==first.lower(),db.func.lower(Member.last_name)==last.lower()).first()
            if duplicate:
                return {"error":"Un membre avec ce prénom et ce nom existe déjà.","duplicateId":duplicate.id},409
            if email and Member.query.filter(db.func.lower(Member.email)==email.lower()).first():
                return {"error":"Un membre avec cet e-mail existe déjà."},409
            member=Member(first_name=first,last_name=last,company=(data.get("company") or "").strip() or None,activity=(data.get("activity") or "").strip() or None,email=email,phone=(data.get("phone") or "").strip() or None)
            db.session.add(member)
            db.session.flush()
            membership=Membership(member_id=member.id,start_date=date.today(),status="ACTIVE")
            db.session.add(membership)
            db.session.commit()
            return {"id":member.id,"firstName":member.first_name,"lastName":member.last_name,"company":member.company,"activity":member.activity,"email":member.email,"phone":member.phone,"status":"ACTIVE"},201
        rows=Member.query.order_by(Member.last_name,Member.first_name).all()
        return jsonify([member_payload(m) for m in rows])

    @app.route("/api/members/<int:member_id>",methods=["GET","PATCH"])
    def member_detail(member_id):
        _,auth_error=require_admin()
        if auth_error:return auth_error
        member=db.session.get(Member,member_id)
        if not member:return {"error":"Membre introuvable."},404
        if request.method=="GET":return member_payload(member)
        data=request.get_json(silent=True) or {}
        for key,attr in [("firstName","first_name"),("lastName","last_name"),("company","company"),("activity","activity"),("phone","phone")]:
            if key in data:setattr(member,attr,(data.get(key) or "").strip() or None)
        if not member.first_name or not member.last_name:return {"error":"Prénom et nom obligatoires."},400
        if "email" in data:
            email=(data.get("email") or "").strip().lower() or None
            other=Member.query.filter(Member.id!=member.id,db.func.lower(Member.email)==email).first() if email else None
            if other:return {"error":"Cette adresse e-mail appartient déjà à un autre membre."},409
            if member.user_account and email and member.user_account.email.lower()!=email:
                used=UserAccount.query.filter(UserAccount.id!=member.user_account.id,db.func.lower(UserAccount.email)==email).first()
                if used:return {"error":"Cette adresse e-mail est déjà liée à un autre compte."},409
                member.user_account.email=email
            member.email=email
        membership=Membership.query.filter_by(member_id=member.id).order_by(Membership.id.desc()).first()
        if not membership:
            membership=Membership(member_id=member.id,status="ACTIVE");db.session.add(membership)
        if "status" in data:membership.status=data["status"]
        if "startDate" in data:membership.start_date=date.fromisoformat(data["startDate"]) if data["startDate"] else None
        if "endDate" in data:membership.end_date=date.fromisoformat(data["endDate"]) if data["endDate"] else None
        db.session.commit()
        return member_payload(member)

    def parse_member_excel(file):
        wb=load_workbook(io.BytesIO(file.read()),read_only=True,data_only=True)
        ws=wb.active
        rows=list(ws.iter_rows(values_only=True))
        if not rows:return [],[],["Le fichier est vide."]
        aliases={"prenom":"firstName","prénom":"firstName","nom":"lastName","societe":"company","société":"company","entreprise":"company","activite":"activity","activité":"activity","email":"email","e-mail":"email","mail":"email","telephone":"phone","téléphone":"phone","gsm":"phone","date entree":"startDate","date d'entrée":"startDate","date entrée":"startDate","statut":"status"}
        import unicodedata
        def norm(v):
            s=str(v or "").strip().lower()
            return " ".join(s.split())
        headers=[norm(v) for v in rows[0]]
        mapping={i:aliases[h] for i,h in enumerate(headers) if h in aliases}
        if "firstName" not in mapping.values() or "lastName" not in mapping.values():
            return [],[],["Colonnes Prénom et Nom obligatoires."]
        preview=[];errors=[]
        for n,row in enumerate(rows[1:],start=2):
            data={}
            for i,key in mapping.items():
                if i>=len(row):continue
                v=row[i]
                if hasattr(v,"isoformat") and key=="startDate":v=v.date().isoformat() if hasattr(v,"date") else v.isoformat()
                data[key]=str(v).strip() if v is not None else ""
            if not data.get("firstName") and not data.get("lastName"):continue
            if not data.get("firstName") or not data.get("lastName"):
                errors.append(f"Ligne {n}: prénom ou nom manquant.");continue
            existing=Member.query.filter(db.func.lower(Member.first_name)==data["firstName"].lower(),db.func.lower(Member.last_name)==data["lastName"].lower()).first()
            data["row"]=n;data["action"]="UPDATE" if existing else "CREATE";data["existingId"]=existing.id if existing else None
            preview.append(data)
        return preview,[{"column":headers[i],"field":key} for i,key in mapping.items()],errors

    @app.post("/api/members/import/preview")
    def import_members_preview():
        _,auth_error=require_admin()
        if auth_error:return auth_error
        file=request.files.get("file")
        if not file or not file.filename.lower().endswith(".xlsx"):return {"error":"Fichier .xlsx requis."},400
        try: preview,mapping,errors=parse_member_excel(file)
        except Exception:return {"error":"Impossible de lire ce fichier Excel."},400
        return {"rows":preview,"mapping":mapping,"errors":errors,"summary":{"total":len(preview),"create":sum(x["action"]=="CREATE" for x in preview),"update":sum(x["action"]=="UPDATE" for x in preview)}}

    @app.post("/api/members/import/confirm")
    def import_members_confirm():
        _,auth_error=require_admin()
        if auth_error:return auth_error
        data=request.get_json(silent=True) or {}
        rows=data.get("rows") or []
        created=updated=0
        for item in rows:
            first=(item.get("firstName") or "").strip();last=(item.get("lastName") or "").strip()
            if not first or not last:continue
            member=Member.query.filter(db.func.lower(Member.first_name)==first.lower(),db.func.lower(Member.last_name)==last.lower()).first()
            if member:updated+=1
            else:
                member=Member(first_name=first,last_name=last);db.session.add(member);db.session.flush();created+=1
            for key,attr in [("company","company"),("activity","activity"),("email","email"),("phone","phone")]:
                if item.get(key):setattr(member,attr,str(item[key]).strip())
            membership=Membership.query.filter_by(member_id=member.id).order_by(Membership.id.desc()).first()
            if not membership:membership=Membership(member_id=member.id,status="ACTIVE");db.session.add(membership)
            if item.get("startDate"):
                try:membership.start_date=date.fromisoformat(str(item["startDate"])[:10])
                except ValueError:pass
            status=(item.get("status") or "").strip().upper()
            if status in {"ACTIVE","ONBOARDING","FORMER"}:membership.status=status
        account,_=bearer_account()
        log=ImportLog(account_id=account.id,filename=(data.get("filename") or "import.xlsx")[:255],created_count=created,updated_count=updated,error_count=0)
        db.session.add(log)
        db.session.commit()
        return {"status":"ok","created":created,"updated":updated}

    @app.get("/api/members/import/history")
    def import_members_history():
        _,auth_error=require_admin()
        if auth_error:return auth_error
        rows=ImportLog.query.order_by(ImportLog.created_at.desc()).limit(50).all()
        return jsonify([{"id":x.id,"filename":x.filename,"created":x.created_count,"updated":x.updated_count,"errors":x.error_count,"createdAt":x.created_at.isoformat(),"by":f"{x.account.member.first_name} {x.account.member.last_name}"} for x in rows])

    @app.get("/api/members/export")
    def export_members():
        _,auth_error=require_admin()
        if auth_error:return auth_error
        wb=Workbook();ws=wb.active;ws.title="Membres"
        ws.append(["Prénom","Nom","Société","Activité","Email","Téléphone","Date entrée","Date sortie","Statut","Accès Hub"])
        for m in Member.query.order_by(Member.last_name,Member.first_name).all():
            membership=Membership.query.filter_by(member_id=m.id).order_by(Membership.id.desc()).first()
            ws.append([m.first_name,m.last_name,m.company or "",m.activity or "",m.email or "",m.phone or "",membership.start_date if membership else None,membership.end_date if membership else None,membership.status if membership else "",m.user_account.status if m.user_account else "NO_ACCOUNT"])
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=min(max(12,max(len(str(cell.value or "")) for cell in col)+2),40)
        out=io.BytesIO();wb.save(out);out.seek(0)
        return send_file(out,as_attachment=True,download_name=f"BNI-Soignies-Membres-{date.today().isoformat()}.xlsx",mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    @app.post("/api/members/<int:member_id>/invite")
    def invite_member(member_id):
        _,auth_error=require_admin()
        if auth_error:return auth_error
        member=db.session.get(Member,member_id)
        if not member:
            return {"error":"Membre introuvable."},404
        if not member.email:
            return {"error":"Une adresse e-mail est nécessaire pour créer l'accès."},400
        existing=UserAccount.query.filter(db.func.lower(UserAccount.email)==member.email.lower()).first()
        if existing and existing.member_id!=member.id:
            return {"error":"Cette adresse e-mail est déjà liée à un autre compte."},409
        account=member.user_account or UserAccount(member_id=member.id,email=member.email.lower())
        account.email=member.email.lower()
        account.status="INVITED"
        account.invitation_token=secrets.token_urlsafe(32)
        account.invitation_expires_at=datetime.utcnow()+timedelta(hours=48)
        if not account.id: db.session.add(account)
        db.session.commit()
        base=os.getenv("APP_URL","https://hub.bim-soignies.be").rstrip("/")
        activation_url=f"{base}/?activate={account.invitation_token}"
        email_sent=False
        host=os.getenv("SMTP_HOST","").strip();sender=os.getenv("SMTP_FROM","").strip()
        if host and sender:
            msg=EmailMessage()
            msg["Subject"]="BNI Soignies Hub — votre invitation"
            msg["From"]=sender;msg["To"]=member.email
            msg.set_content(f"""Bonjour {member.first_name},

Vous êtes invité(e) à découvrir BNI Soignies Hub.

Xavier vous propose de tester l'application et de lui remonter vos remarques, idées ou éventuels problèmes rencontrés.

Activez votre accès avec ce lien (valable 48 heures) :
{activation_url}

À bientôt sur BNI Soignies Hub.""")
            port=int(os.getenv("SMTP_PORT","587"));user=os.getenv("SMTP_USER","").strip();password=os.getenv("SMTP_PASSWORD","")
            use_ssl=os.getenv("SMTP_SSL","false").lower() in {"1","true","yes"}
            smtp=(smtplib.SMTP_SSL(host,port,timeout=15) if use_ssl else smtplib.SMTP(host,port,timeout=15))
            try:
                if not use_ssl and os.getenv("SMTP_STARTTLS","true").lower() in {"1","true","yes"}:smtp.starttls()
                if user:smtp.login(user,password)
                smtp.send_message(msg);email_sent=True
            except Exception as exc:
                app.logger.exception("Invitation email delivery failed for member %s: %s",member.id,exc)
            finally:
                try:smtp.quit()
                except Exception:pass
        return {"status":"INVITED","activationUrl":activation_url,"expiresInHours":48,"emailSent":email_sent}

    @app.post("/api/auth/activate")
    def activate():
        data=request.get_json(silent=True) or {}
        token=(data.get("token") or "").strip()
        password=data.get("password") or ""
        if len(password)<10:
            return {"error":"Le mot de passe doit contenir au moins 10 caractères."},400
        account=UserAccount.query.filter_by(invitation_token=token).first()
        if not account or not account.invitation_expires_at or account.invitation_expires_at<datetime.utcnow():
            return {"error":"Lien d'activation invalide ou expiré."},400
        account.password_hash=generate_password_hash(password)
        account.status="ACTIVE"
        account.activated_at=datetime.utcnow()
        account.invitation_token=None
        account.invitation_expires_at=None
        db.session.commit()
        return {"status":"ACTIVE","message":"Compte activé."}

    def account_payload(account):
        today=date.today()
        active_mandates=Mandate.query.filter(Mandate.start_date<=today,Mandate.end_date>=today).all()
        mandate_ids=[m.id for m in active_mandates]
        roles=[]
        if mandate_ids:
            assignments=RoleAssignment.query.filter(RoleAssignment.member_id==account.member_id,RoleAssignment.mandate_id.in_(mandate_ids),RoleAssignment.status=="ACTIVE").all()
            role_ids=[a.role_id for a in assignments]
            if role_ids: roles=[r.name for r in Role.query.filter(Role.id.in_(role_ids)).all()]
        admin=bool(account.is_system_admin or any(r.lower() in {"président","president","vice-président","vice-president","secrétaire-trésorier","secretaire-tresorier"} for r in roles))
        return {"id":account.member.id,"firstName":account.member.first_name,"lastName":account.member.last_name,"email":account.email,"roles":roles,"isAdmin":admin}

    def bearer_account():
        header=request.headers.get("Authorization","")
        if not header.startswith("Bearer "): return None,None
        raw=header[7:].strip()
        import hashlib
        hashed=hashlib.sha256(raw.encode()).hexdigest()
        session=AuthSession.query.filter_by(token_hash=hashed).first()
        if not session or session.expires_at<datetime.utcnow():
            if session:
                db.session.delete(session);db.session.commit()
            return None,None
        return session.account,session

    @app.post("/api/auth/login")
    def login():
        data=request.get_json(silent=True) or {}
        email=(data.get("email") or "").strip().lower()
        account=UserAccount.query.filter(db.func.lower(UserAccount.email)==email).first()
        if not account or account.status!="ACTIVE" or not account.password_hash or not check_password_hash(account.password_hash,data.get("password") or ""):
            return {"error":"Identifiants incorrects."},401
        raw=secrets.token_urlsafe(48)
        import hashlib
        session=AuthSession(account_id=account.id,token_hash=hashlib.sha256(raw.encode()).hexdigest(),expires_at=datetime.utcnow()+timedelta(days=7))
        db.session.add(session);db.session.commit()
        return {"token":raw,"expiresInDays":7,"member":account_payload(account)}

    @app.get("/api/auth/me")
    def me():
        account,_=bearer_account()
        if not account:return {"error":"Session invalide ou expirée."},401
        return {"member":account_payload(account)}

    @app.post("/api/auth/logout")
    def logout():
        _,session=bearer_account()
        if session:
            db.session.delete(session);db.session.commit()
        return {"status":"ok"}

    @app.post("/api/system/bootstrap-xavier")
    def bootstrap_xavier():
        key=request.headers.get("X-Bootstrap-Key","")
        if not key or not secrets.compare_digest(key,os.getenv("BOOTSTRAP_KEY","disabled")):
            return {"error":"Forbidden"},403
        candidates=Member.query.filter(db.func.lower(Member.first_name)=="xavier",db.func.lower(Member.last_name)=="renaud").all()
        active=None
        for m in candidates:
            if m.user_account and m.user_account.status=="ACTIVE" and m.user_account.password_hash:
                active=m;break
        if not active:
            return {"error":"Aucun compte Xavier Renaud actif trouvé."},404
        active.user_account.is_system_admin=True
        membership=Membership.query.filter_by(member_id=active.id).order_by(Membership.id.asc()).first()
        if not membership:
            membership=Membership(member_id=active.id,status="ACTIVE");db.session.add(membership)
        membership.start_date=date(2017,4,1)
        membership.end_date=None
        membership.status="ACTIVE"
        mandate=Mandate.query.filter_by(name="Octobre 2026 - Mars 2027").first()
        if not mandate:
            mandate=Mandate(name="Octobre 2026 - Mars 2027",start_date=date(2026,10,1),end_date=date(2027,3,31),status="PREPARATION");db.session.add(mandate);db.session.flush()
        role=Role.query.filter(db.func.lower(Role.name)=="président").first()
        if not role:
            role=Role(name="Président");db.session.add(role);db.session.flush()
        assignment=RoleAssignment.query.filter_by(mandate_id=mandate.id,role_id=role.id,member_id=active.id).first()
        if not assignment:
            assignment=RoleAssignment(mandate_id=mandate.id,role_id=role.id,member_id=active.id,status="ACTIVE");db.session.add(assignment)
        db.session.commit()
        return {"status":"ok","memberId":active.id,"systemAdmin":True,"membershipStart":"2017-04-01","mandate":mandate.name,"role":"Président","roleStarts":"2026-10-01","duplicateCandidates":len(candidates)}

    def mandate_payload(m):
        assignments=RoleAssignment.query.filter_by(mandate_id=m.id,status="ACTIVE").all()
        role_ids=list({a.role_id for a in assignments})
        roles={r.id:r for r in Role.query.filter(Role.id.in_(role_ids)).all()} if role_ids else {}
        member_ids=list({a.member_id for a in assignments})
        members={x.id:x for x in Member.query.filter(Member.id.in_(member_ids)).all()} if member_ids else {}
        return {"id":m.id,"name":m.name,"startDate":m.start_date.isoformat(),"endDate":m.end_date.isoformat(),"status":m.status,
                "committee":[{"assignmentId":a.id,"roleId":a.role_id,"role":roles[a.role_id].name if a.role_id in roles else "—","memberId":a.member_id,
                "member":f"{members[a.member_id].first_name} {members[a.member_id].last_name}" if a.member_id in members else "—"} for a in assignments]}

    def handover_payload(x):
        role=db.session.get(Role,x.role_id);fm=db.session.get(Mandate,x.from_mandate_id);tm=db.session.get(Mandate,x.to_mandate_id)
        return {"id":x.id,"title":x.title,"notes":x.notes,"roleId":x.role_id,"role":role.name if role else "—","fromMandate":fm.name if fm else "—","toMandate":tm.name if tm else "—","outgoingConfirmed":bool(x.outgoing_confirmed_at),"incomingConfirmed":bool(x.incoming_confirmed_at),"complete":bool(x.outgoing_confirmed_at and x.incoming_confirmed_at)}

    @app.post("/api/handovers/generate")
    def handovers_generate():
        _,err=require_admin()
        if err:return err
        d=request.get_json(silent=True) or {}
        try:from_id=int(d["fromMandateId"]);to_id=int(d["toMandateId"])
        except (KeyError,TypeError,ValueError):return {"error":"Mandatures obligatoires."},400
        incoming=RoleAssignment.query.filter_by(mandate_id=to_id,status="ACTIVE").all()
        created=0
        defaults=["Points clés et dossiers en cours","Procédures et échéances","Ressources et documents","Plateformes, accès et renouvellements"]
        for a in incoming:
            outgoing=RoleAssignment.query.filter_by(mandate_id=from_id,role_id=a.role_id,status="ACTIVE").first()
            if not outgoing:continue
            for title in defaults:
                exists=HandoverItem.query.filter_by(from_mandate_id=from_id,to_mandate_id=to_id,role_id=a.role_id,title=title).first()
                if not exists:
                    db.session.add(HandoverItem(from_mandate_id=from_id,to_mandate_id=to_id,role_id=a.role_id,title=title));created+=1
        db.session.commit()
        return {"created":created,"items":[handover_payload(x) for x in HandoverItem.query.filter_by(from_mandate_id=from_id,to_mandate_id=to_id).order_by(HandoverItem.role_id,HandoverItem.id).all()]}

    @app.get("/api/handovers")
    def handovers_list():
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        if account.is_system_admin:return jsonify([handover_payload(x) for x in HandoverItem.query.order_by(HandoverItem.id.desc()).all()])
        role_ids=current_role_ids(account)
        rows=HandoverItem.query.filter(HandoverItem.role_id.in_(role_ids)).order_by(HandoverItem.id.desc()).all() if role_ids else []
        return jsonify([handover_payload(x) for x in rows])

    @app.get("/api/me/handovers")
    def my_handovers():
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        rows=HandoverItem.query.order_by(HandoverItem.id.desc()).all()
        out=[]
        for x in rows:
            outgoing=RoleAssignment.query.filter_by(mandate_id=x.from_mandate_id,role_id=x.role_id,member_id=account.member_id,status="ACTIVE").first()
            incoming=RoleAssignment.query.filter_by(mandate_id=x.to_mandate_id,role_id=x.role_id,member_id=account.member_id,status="ACTIVE").first()
            if outgoing or incoming:
                p=handover_payload(x);p["mySide"]="outgoing" if outgoing else "incoming";out.append(p)
        return jsonify(out)

    @app.post("/api/handovers")
    def handovers_create():
        _,err=require_admin()
        if err:return err
        d=request.get_json(silent=True) or {}
        try:x=HandoverItem(from_mandate_id=int(d["fromMandateId"]),to_mandate_id=int(d["toMandateId"]),role_id=int(d["roleId"]),title=d["title"].strip(),notes=d.get("notes"))
        except (KeyError,TypeError,ValueError):return {"error":"Données de transmission incomplètes."},400
        db.session.add(x);db.session.commit();return handover_payload(x),201

    @app.patch("/api/handovers/<int:item_id>/confirm")
    def handover_confirm(item_id):
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        x=db.session.get(HandoverItem,item_id)
        if not x:return {"error":"Transmission introuvable."},404
        d=request.get_json(silent=True) or {};side=d.get("side")
        if account.is_system_admin:
            allowed=True
        else:
            mandate_id=x.from_mandate_id if side=="outgoing" else x.to_mandate_id
            allowed=RoleAssignment.query.filter_by(mandate_id=mandate_id,role_id=x.role_id,member_id=account.member_id,status="ACTIVE").first() is not None
        if not allowed:return {"error":"Cette validation appartient au responsable de la fonction."},403
        if side=="outgoing":x.outgoing_confirmed_at=datetime.utcnow()
        elif side=="incoming":x.incoming_confirmed_at=datetime.utcnow()
        else:return {"error":"Sens de validation invalide."},400
        audit(account,"HANDOVER_CONFIRMED","handover",x.id,f"side={side}")
        db.session.commit();return handover_payload(x)

    def board_member(account):
        if account.is_system_admin:return True
        today=date.today();mandate=Mandate.query.filter(Mandate.start_date<=today,Mandate.end_date>=today).order_by(Mandate.start_date.desc()).first()
        if not mandate:return False
        allowed=["président","president","vice-président","vice president","vice-president","secrétaire trésorier","secretaire tresorier","secrétaire-trésorier","secretaire-tresorier"]
        assignments=RoleAssignment.query.filter_by(mandate_id=mandate.id,member_id=account.member_id,status="ACTIVE").all()
        return any((db.session.get(Role,x.role_id) and db.session.get(Role,x.role_id).name.lower().strip() in allowed) for x in assignments)

    def current_role_ids(account):
        today=date.today();mandates=Mandate.query.filter(Mandate.start_date<=today,Mandate.end_date>=today).all()
        mids=[m.id for m in mandates]
        if not mids:return set()
        return {x.role_id for x in RoleAssignment.query.filter(RoleAssignment.mandate_id.in_(mids),RoleAssignment.member_id==account.member_id,RoleAssignment.status=="ACTIVE").all()}

    def can_view_library_item(account,role_id,board_only):
        if account.is_system_admin:return True
        if board_only and not board_member(account):return False
        return not role_id or role_id in current_role_ids(account)

    ALLOWED_RESOURCE_EXTENSIONS={"pdf","doc","docx","xls","xlsx","png","jpg","jpeg","webp"}
    def resource_meta(x):
        return {"id":x.id,"title":x.title,"description":x.description,"category":x.category,"url":x.url,"roleId":x.role_id,"boardOnly":x.board_only,"fileName":x.file_name,"hasFile":bool(x.file_data)}

    @app.post("/api/resources/upload")
    def resources_upload():
        _,err=require_admin()
        if err:return err
        file=request.files.get("file")
        if not file or not file.filename:return {"error":"Fichier obligatoire."},400
        ext=file.filename.rsplit(".",1)[-1].lower() if "." in file.filename else ""
        if ext not in ALLOWED_RESOURCE_EXTENSIONS:return {"error":"Type de fichier non autorisé."},400
        data=file.read()
        if len(data)>15*1024*1024:return {"error":"Fichier trop volumineux (15 Mo maximum)."},413
        x=Resource(title=(request.form.get("title") or file.filename).strip(),description=request.form.get("description"),category=request.form.get("category") or "GENERAL",role_id=int(request.form["roleId"]) if request.form.get("roleId") else None,board_only=request.form.get("boardOnly")=="true",file_name=file.filename,file_type=file.mimetype,file_data=data)
        db.session.add(x);db.session.flush();audit(_, "RESOURCE_FILE_UPLOADED","resource",x.id,x.file_name);db.session.commit();return resource_meta(x),201

    @app.get("/api/resources/<int:item_id>/file")
    def resource_file(item_id):
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        x=db.session.get(Resource,item_id)
        if not x or not x.active or not x.file_data:return {"error":"Fichier introuvable."},404
        if not can_view_library_item(account,x.role_id,x.board_only):return {"error":"Accès refusé."},403
        return send_file(io.BytesIO(x.file_data),mimetype=x.file_type or "application/octet-stream",download_name=x.file_name or "ressource",as_attachment=True)

    @app.get("/api/resources")
    def resources_list():
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        rows=Resource.query.filter_by(active=True).order_by(Resource.category,Resource.title).all()
        return jsonify([resource_meta(x) for x in rows if can_view_library_item(account,x.role_id,x.board_only)])

    @app.post("/api/resources")
    def resources_create():
        _,err=require_admin()
        if err:return err
        d=request.get_json(silent=True) or {}
        if not d.get("title"):return {"error":"Titre obligatoire."},400
        x=Resource(title=d["title"].strip(),description=d.get("description"),category=d.get("category","GENERAL"),url=d.get("url"),role_id=d.get("roleId") or None,board_only=bool(d.get("boardOnly",False)))
        db.session.add(x);db.session.commit();return {"id":x.id},201

    @app.patch("/api/resources/<int:item_id>")
    def resources_update(item_id):
        _,err=require_admin()
        if err:return err
        x=db.session.get(Resource,item_id)
        if not x:return {"error":"Ressource introuvable."},404
        d=request.get_json(silent=True) or {}
        for key,attr in [("title","title"),("description","description"),("category","category"),("url","url")]:
            if key in d:setattr(x,attr,d[key])
        if "roleId" in d:x.role_id=d.get("roleId") or None
        if "boardOnly" in d:x.board_only=bool(d["boardOnly"])
        if "active" in d:x.active=bool(d["active"])
        db.session.commit();return {"status":"ok"}

    @app.delete("/api/resources/<int:item_id>")
    def resources_delete(item_id):
        _,err=require_admin()
        if err:return err
        x=db.session.get(Resource,item_id)
        if not x:return {"error":"Ressource introuvable."},404
        x.active=False;audit(_, "RESOURCE_ARCHIVED","resource",x.id,x.title);db.session.commit();return {"status":"archived"}

    @app.get("/api/platforms")
    def platforms_list():
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        rows=Platform.query.filter_by(active=True).order_by(Platform.name).all()
        return jsonify([{"id":x.id,"name":x.name,"description":x.description,"url":x.url,"ownerRoleId":x.owner_role_id,"boardOnly":x.board_only,"renewalInfo":x.renewal_info} for x in rows if can_view_library_item(account,x.owner_role_id,x.board_only)])

    @app.post("/api/platforms")
    def platforms_create():
        _,err=require_admin()
        if err:return err
        d=request.get_json(silent=True) or {}
        if not d.get("name"):return {"error":"Nom obligatoire."},400
        x=Platform(name=d["name"].strip(),description=d.get("description"),url=d.get("url"),owner_role_id=d.get("ownerRoleId") or None,board_only=bool(d.get("boardOnly",False)),renewal_info=d.get("renewalInfo"))
        db.session.add(x);db.session.flush();audit(_, "PLATFORM_CREATED","platform",x.id,x.name);db.session.commit();return {"id":x.id},201

    @app.patch("/api/platforms/<int:item_id>")
    def platforms_update(item_id):
        _,err=require_admin()
        if err:return err
        x=db.session.get(Platform,item_id)
        if not x:return {"error":"Plateforme introuvable."},404
        d=request.get_json(silent=True) or {}
        for key,attr in [("name","name"),("description","description"),("url","url"),("renewalInfo","renewal_info")]:
            if key in d:setattr(x,attr,d[key])
        if "ownerRoleId" in d:x.owner_role_id=d.get("ownerRoleId") or None
        if "boardOnly" in d:x.board_only=bool(d["boardOnly"])
        if "active" in d:x.active=bool(d["active"])
        audit(_, "PLATFORM_UPDATED","platform",x.id,x.name);db.session.commit();return {"status":"ok"}

    @app.delete("/api/platforms/<int:item_id>")
    def platforms_delete(item_id):
        _,err=require_admin()
        if err:return err
        x=db.session.get(Platform,item_id)
        if not x:return {"error":"Plateforme introuvable."},404
        x.active=False;audit(_, "PLATFORM_ARCHIVED","platform",x.id,x.name);db.session.commit();return {"status":"archived"}

    def require_alumni_access():
        account,_=bearer_account()
        if not account:return None,({"error":"Authentification requise."},401)
        if account.is_system_admin:return account,None
        today=date.today()
        mandate=Mandate.query.filter(Mandate.start_date<=today,Mandate.end_date>=today).order_by(Mandate.start_date.desc()).first()
        if not mandate:return account,({"error":"Aucune mandature active."},403)
        allowed=["président","president","vice-président","vice president","vice-president","secrétaire trésorier","secretaire tresorier","secrétaire-trésorier","secretaire-tresorier"]
        assignments=RoleAssignment.query.filter_by(mandate_id=mandate.id,member_id=account.member_id,status="ACTIVE").all()
        roles=[db.session.get(Role,x.role_id) for x in assignments]
        if any(r and r.name.lower().strip() in allowed for r in roles):return account,None
        return account,({"error":"Accès réservé au Board de la mandature active."},403)

    @app.get("/api/alumni")
    def alumni_list():
        _,err=require_alumni_access()
        if err:return err
        rows=db.session.query(Member,Membership).join(Membership,Membership.member_id==Member.id).filter(Membership.status.in_(["FORMER","ALUMNI"])).order_by(Member.last_name,Member.first_name).all()
        return jsonify([{"memberId":m.id,"firstName":m.first_name,"lastName":m.last_name,"company":m.company,"activity":m.activity,"email":m.email,"phone":m.phone,"startDate":ms.start_date.isoformat() if ms.start_date else None,"endDate":ms.end_date.isoformat() if ms.end_date else None} for m,ms in rows])

    @app.get("/api/alumni/export")
    def alumni_export():
        _,err=require_alumni_access()
        if err:return err
        from io import BytesIO
        from openpyxl import Workbook
        wb=Workbook();ws=wb.active;ws.title="BNI Alumni"
        ws.append(["Prénom","Nom","Société","Activité","Email","Téléphone","Date entrée","Date sortie"])
        rows=db.session.query(Member,Membership).join(Membership,Membership.member_id==Member.id).filter(Membership.status.in_(["FORMER","ALUMNI"])).order_by(Member.last_name,Member.first_name).all()
        for m,ms in rows:ws.append([m.first_name,m.last_name,m.company,m.activity,m.email,m.phone,ms.start_date,ms.end_date])
        out=BytesIO();wb.save(out);out.seek(0)
        return send_file(out,as_attachment=True,download_name="BNI_Soignies_Alumni.xlsx",mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    ONBOARDING_DEFAULTS=[
        ("OBSERVE_SETUP","Observer la préparation de salle"),
        ("OBSERVE_RECEPTION","Observer l'accueil des invités"),
        ("OBSERVE_FOLLOWUP","Observer le suivi des invités"),
        ("INFOMERCIAL","Formation infomercial"),
        ("INVITATION","Formation invitation"),
        ("MENTOR_VALIDATION","Validation finale par le mentor"),
    ]

    def onboarding_payload(member_id):
        m=db.session.get(Member,member_id)
        membership=Membership.query.filter_by(member_id=member_id).order_by(Membership.id.desc()).first()
        mentorship=Mentorship.query.filter_by(member_id=member_id,status="ACTIVE").order_by(Mentorship.id.desc()).first()
        mentor=db.session.get(Member,mentorship.mentor_id) if mentorship else None
        steps=OnboardingStep.query.filter_by(member_id=member_id).order_by(OnboardingStep.id).all()
        return {"memberId":member_id,"member":f"{m.first_name} {m.last_name}" if m else "—","status":membership.status if membership else None,
          "mentorId":mentor.id if mentor else None,"mentor":f"{mentor.first_name} {mentor.last_name}" if mentor else None,
          "steps":[{"id":x.id,"code":x.code,"label":x.label,"completed":x.completed,"completedAt":x.completed_at.isoformat() if x.completed_at else None} for x in steps],
          "progress":round(100*sum(x.completed for x in steps)/len(steps)) if steps else 0}

    @app.get("/api/onboarding")
    def onboarding_list():
        _,auth_error=require_admin()
        if auth_error:return auth_error
        ids=[x.member_id for x in Membership.query.filter(Membership.status=="ONBOARDING").all()]
        return jsonify([onboarding_payload(i) for i in ids])

    @app.post("/api/members/<int:member_id>/onboarding")
    def start_onboarding(member_id):
        _,auth_error=require_admin()
        if auth_error:return auth_error
        data=request.get_json(silent=True) or {}
        try:mentor_id=int(data.get("mentorId"))
        except (TypeError,ValueError):return {"error":"Mentor obligatoire."},400
        if mentor_id==member_id or not db.session.get(Member,mentor_id):return {"error":"Mentor invalide."},400
        membership=Membership.query.filter_by(member_id=member_id).order_by(Membership.id.desc()).first()
        if not membership:return {"error":"Adhésion introuvable."},404
        membership.status="ONBOARDING"
        Mentorship.query.filter_by(member_id=member_id,status="ACTIVE").update({"status":"COMPLETED","completed_at":date.today()})
        db.session.add(Mentorship(member_id=member_id,mentor_id=mentor_id,status="ACTIVE",started_at=date.today()))
        existing={x.code for x in OnboardingStep.query.filter_by(member_id=member_id).all()}
        for code,label in ONBOARDING_DEFAULTS:
            if code not in existing:db.session.add(OnboardingStep(member_id=member_id,code=code,label=label))
        db.session.commit();return onboarding_payload(member_id),201

    @app.patch("/api/onboarding/<int:member_id>/steps/<int:step_id>")
    def onboarding_step(member_id,step_id):
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        step=db.session.get(OnboardingStep,step_id)
        if not step or step.member_id!=member_id:return {"error":"Étape introuvable."},404
        mentorship=Mentorship.query.filter_by(member_id=member_id,status="ACTIVE").first()
        if not (account.is_system_admin or (mentorship and mentorship.mentor_id==account.member_id)):return {"error":"Validation réservée au mentor ou à l'administrateur."},403
        completed=bool((request.get_json(silent=True) or {}).get("completed",True))
        step.completed=completed;step.completed_at=datetime.utcnow() if completed else None
        db.session.commit();return onboarding_payload(member_id)

    @app.post("/api/onboarding/<int:member_id>/complete")
    def complete_onboarding(member_id):
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        mentorship=Mentorship.query.filter_by(member_id=member_id,status="ACTIVE").first()
        if not (account.is_system_admin or (mentorship and mentorship.mentor_id==account.member_id)):return {"error":"Validation réservée au mentor ou à l'administrateur."},403
        steps=OnboardingStep.query.filter_by(member_id=member_id).all()
        if not steps or not all(x.completed for x in steps):return {"error":"Toutes les étapes doivent être validées."},400
        membership=Membership.query.filter_by(member_id=member_id).order_by(Membership.id.desc()).first()
        membership.status="ACTIVE"
        if mentorship:mentorship.status="COMPLETED";mentorship.completed_at=date.today()
        audit(account,"ONBOARDING_COMPLETED","member",member_id)
        notify(member_id,"ONBOARDING_COMPLETE","Intégration terminée","Votre mentor a validé votre intégration. Vous êtes maintenant membre autonome.","profile")
        db.session.commit();return onboarding_payload(member_id)

    @app.get("/api/audit")
    def audit_list():
        account,err=require_admin()
        if err:return err
        q=AuditLog.query.order_by(AuditLog.created_at.desc()).limit(300).all()
        out=[]
        for x in q:
            m=db.session.get(Member,x.member_id) if x.member_id else None
            out.append({"id":x.id,"actor":f"{m.first_name} {m.last_name}" if m else "Système","action":x.action,"entityType":x.entity_type,"entityId":x.entity_id,"details":x.details,"createdAt":x.created_at.isoformat()})
        return jsonify(out)

    @app.get("/api/mandates/readiness")
    def mandate_readiness():
        account,err=require_admin()
        if err:return err
        today=date.today();nxt=Mandate.query.filter(Mandate.start_date>today).order_by(Mandate.start_date.asc()).first()
        if not nxt:return {"ready":False,"blocking":["Aucune prochaine mandature n'est configurée."],"warnings":[]},200
        assignments=RoleAssignment.query.filter_by(mandate_id=nxt.id,status="ACTIVE").all()
        role_names={r.id:r.name for r in Role.query.all()}
        names=[role_names.get(a.role_id,"").lower() for a in assignments]
        blocking=[];warnings=[]
        if not any(n in {"président","president"} for n in names):blocking.append("Aucun Président n'est désigné pour la prochaine mandature.")
        if not assignments:blocking.append("Le futur comité ne contient encore aucune affectation.")
        duplicate_members={}
        for a in assignments:duplicate_members.setdefault(a.member_id,[]).append(a.role_id)
        multi=[mid for mid,rs in duplicate_members.items() if len(rs)>1]
        if multi:warnings.append(f"{len(multi)} membre(s) cumulent plusieurs fonctions dans la prochaine mandature.")
        handovers=HandoverItem.query.filter_by(to_mandate_id=nxt.id).all()
        incomplete=[x for x in handovers if not (x.outgoing_confirmed_at and x.incoming_confirmed_at)]
        if not handovers:warnings.append("Aucune checklist de passage de flambeau n'a encore été générée.")
        elif incomplete:warnings.append(f"{len(incomplete)} élément(s) de transmission restent à valider.")
        if nxt.start_date>today:
            days=(nxt.start_date-today).days
        else:days=0
        return {"ready":not blocking,"nextMandate":{"id":nxt.id,"name":nxt.name,"startDate":nxt.start_date.isoformat(),"endDate":nxt.end_date.isoformat()},"daysUntilStart":days,"assignments":len(assignments),"handoverTotal":len(handovers),"handoverIncomplete":len(incomplete),"blocking":blocking,"warnings":warnings,"systemAdminProtected":bool(account.is_system_admin)}

    @app.get("/api/mandates/status")
    def mandate_status():
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        today=date.today();current=active_mandate(today)
        nxt=Mandate.query.filter(Mandate.start_date>today).order_by(Mandate.start_date.asc()).first()
        previous=Mandate.query.filter(Mandate.end_date<today).order_by(Mandate.end_date.desc()).first()
        def p(x):return None if not x else {"id":x.id,"name":x.name,"startDate":x.start_date.isoformat(),"endDate":x.end_date.isoformat()}
        return {"today":today.isoformat(),"current":p(current),"next":p(nxt),"previous":p(previous),"systemAdmin":bool(account.is_system_admin)}

    @app.get("/api/mandates")
    def mandates_list():
        _,auth_error=require_admin()
        if auth_error:return auth_error
        rows=Mandate.query.order_by(Mandate.start_date.desc()).all()
        return jsonify([mandate_payload(x) for x in rows])

    @app.route("/api/roles",methods=["GET","POST"])
    def roles_api():
        _,auth_error=require_admin()
        if auth_error:return auth_error
        if request.method=="POST":
            data=request.get_json(silent=True) or {};name=(data.get("name") or "").strip()
            if not name:return {"error":"Nom de fonction obligatoire."},400
            existing=Role.query.filter(db.func.lower(Role.name)==name.lower()).first()
            if existing:return {"id":existing.id,"name":existing.name}
            role=Role(name=name);db.session.add(role);db.session.commit()
            return {"id":role.id,"name":role.name},201
        return jsonify([{"id":r.id,"name":r.name} for r in Role.query.order_by(Role.name).all()])

    @app.post("/api/mandates/<int:mandate_id>/committee")
    def committee_add(mandate_id):
        _,auth_error=require_admin()
        if auth_error:return auth_error
        mandate=db.session.get(Mandate,mandate_id)
        if not mandate:return {"error":"Mandat introuvable."},404
        data=request.get_json(silent=True) or {}
        try:member_id=int(data.get("memberId"));role_id=int(data.get("roleId"))
        except (TypeError,ValueError):return {"error":"Membre et fonction obligatoires."},400
        if not db.session.get(Member,member_id) or not db.session.get(Role,role_id):return {"error":"Membre ou fonction introuvable."},404
        assignment=RoleAssignment.query.filter_by(mandate_id=mandate_id,member_id=member_id,role_id=role_id).first()
        if assignment:assignment.status="ACTIVE"
        else:assignment=RoleAssignment(mandate_id=mandate_id,member_id=member_id,role_id=role_id,status="ACTIVE");db.session.add(assignment)
        db.session.commit();return mandate_payload(mandate),201

    @app.delete("/api/mandates/<int:mandate_id>/committee/<int:assignment_id>")
    def committee_remove(mandate_id,assignment_id):
        _,auth_error=require_admin()
        if auth_error:return auth_error
        assignment=db.session.get(RoleAssignment,assignment_id)
        if not assignment or assignment.mandate_id!=mandate_id:return {"error":"Attribution introuvable."},404
        assignment.status="INACTIVE";db.session.commit();return {"status":"ok"}

    def meeting_payload(m):
        return {"id":m.id,"date":m.date.isoformat(),"startTime":m.start_time.strftime("%H:%M") if m.start_time else None,"type":m.type,"location":m.location,"status":m.status,"notes":m.notes}

    def assignment_payload(a):
        duty=db.session.get(Duty,a.duty_id);member=db.session.get(Member,a.member_id);meeting=db.session.get(Meeting,a.meeting_id)
        return {"id":a.id,"meetingId":a.meeting_id,"date":meeting.date.isoformat() if meeting else None,"dutyId":a.duty_id,"duty":duty.name if duty else "—","memberId":a.member_id,"member":f"{member.first_name} {member.last_name}" if member else "—","status":a.status}

    @app.get("/api/duties")
    def duties_list():
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        return jsonify([{"id":d.id,"code":d.code,"name":d.name,"active":d.active} for d in Duty.query.order_by(Duty.id).all()])

    @app.get("/api/duty-assignments")
    def duty_assignments_list():
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        rows=DutyAssignment.query.join(Meeting,DutyAssignment.meeting_id==Meeting.id).order_by(Meeting.date,DutyAssignment.duty_id).all()
        return jsonify([assignment_payload(a) for a in rows])

    REQUEST_ROLE_HINTS={
        "INFOMERCIAL":["éducation","education","formation"],
        "COMMUNICATION":["réseaux sociaux","reseaux sociaux","communication"],
        "TRAINING":["éducation","education","formation"],
        "MENTORING":["mentor","mentorat"],
        "EVENT":["événement","evenement","events"],
        "INVITATION":["accueil","invité","invite"],
    }

    def request_assignees(kind,when=None):
        # Explicit admin routing takes priority over mandate-based automatic routing.
        configured=RequestRecipient.query.filter_by(request_type=kind).order_by(RequestRecipient.id).all()
        if configured:
            mids=list(dict.fromkeys([x.member_id for x in configured]))
            members={m.id:m for m in Member.query.filter(Member.id.in_(mids)).all()}
            return [{"memberId":mid,"member":f"{members[mid].first_name} {members[mid].last_name}","email":members[mid].email,"role":"Destinataire configuré"} for mid in mids if mid in members]
        when=when or date.today()
        mandates=Mandate.query.filter(Mandate.start_date<=when,Mandate.end_date>=when).all()
        if not mandates:return []
        hints=REQUEST_ROLE_HINTS.get(kind,[])
        roles=Role.query.all()
        role_ids=[r.id for r in roles if any(h in r.name.lower() for h in hints)]
        if not role_ids:return []
        rows=RoleAssignment.query.filter(RoleAssignment.mandate_id.in_([m.id for m in mandates]),RoleAssignment.role_id.in_(role_ids),RoleAssignment.status=="ACTIVE").all()
        members={m.id:m for m in Member.query.filter(Member.id.in_(list({a.member_id for a in rows}))).all()} if rows else {}
        rolemap={r.id:r.name for r in roles}
        return [{"memberId":a.member_id,"member":f"{members[a.member_id].first_name} {members[a.member_id].last_name}","email":members[a.member_id].email,"role":rolemap.get(a.role_id,"—")} for a in rows if a.member_id in members]

    def can_manage_request(account,kind):
        if account.is_system_admin:return True
        return any(x["memberId"]==account.member_id for x in request_assignees(kind))

    def member_request_payload(x):
        m=db.session.get(Member,x.member_id);assignees=request_assignees(x.type,x.created_at.date() if x.created_at else None)
        return {"id":x.id,"type":x.type,"title":x.title,"message":x.message,"guestName":x.guest_name,"guestEmail":x.guest_email,"eventDate":x.event_date.isoformat() if x.event_date else None,"status":x.status,"createdAt":x.created_at.isoformat(),"member":f"{m.first_name} {m.last_name}" if m else "—","assignees":[{k:v for k,v in t.items() if k!="email"} for t in assignees],"routingStatus":"ROUTED" if assignees else "UNASSIGNED","attachmentName":x.attachment_name,"hasAttachment":bool(x.attachment_data)}

    def notify(member_id,kind,title,message=None,link=None):
        n=Notification(member_id=member_id,type=kind,title=title,message=message,link=link)
        db.session.add(n);return n

    def notification_payload(n):
        return {"id":n.id,"type":n.type,"title":n.title,"message":n.message,"link":n.link,"read":bool(n.read_at),"createdAt":n.created_at.isoformat()}

    @app.get("/api/me/notifications")
    def my_notifications():
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        rows=Notification.query.filter_by(member_id=account.member_id).order_by(Notification.created_at.desc()).limit(100).all()
        return {"unread":sum(not x.read_at for x in rows),"items":[notification_payload(x) for x in rows]}

    @app.post("/api/me/notifications/<int:notification_id>/read")
    def read_notification(notification_id):
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        n=db.session.get(Notification,notification_id)
        if not n or n.member_id!=account.member_id:return {"error":"Notification introuvable."},404
        n.read_at=datetime.utcnow();db.session.commit();return notification_payload(n)

    @app.post("/api/me/notifications/read-all")
    def read_all_notifications():
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        Notification.query.filter_by(member_id=account.member_id,read_at=None).update({"read_at":datetime.utcnow()})
        db.session.commit();return {"status":"ok"}

    REQUEST_ALLOWED_EXTENSIONS={"pdf","png","jpg","jpeg","webp","ppt","pptx","mp4","mov","m4v"}
    REQUEST_MAX_FILE_SIZE=5*1024*1024

    def send_request_email(x,account,targets):
        host=os.getenv("SMTP_HOST","").strip()
        sender=os.getenv("SMTP_FROM","").strip()
        if not host or not sender:return False
        recipients=list(dict.fromkeys([t.get("email") for t in targets if t.get("email")]))
        # Safety net: never silently drop an email when a request has no routed
        # responsible member (or the responsible member has no email yet).
        # The Hub keeps the request internally and sends it to the configured
        # fallback/admin mailbox so it can still be handled.
        if not recipients:
            fallback=(os.getenv("REQUEST_FALLBACK_EMAIL","").strip()
                      or os.getenv("SYSTEM_ADMIN_EMAIL","").strip()
                      or sender)
            if fallback:
                recipients=[fallback]
                app.logger.warning("Member request %s (%s) has no emailed assignee; using fallback recipient.",x.id,x.type)
        if not recipients:return False
        msg=EmailMessage()
        labels={"INFOMERCIAL":"Infomercial","COMMUNICATION":"Communication","TRAINING":"Formation","MENTORING":"Mentorat","EVENT":"Événement","INVITATION":"Invitation"}
        member=account.member
        msg["Subject"]=f"BNI Soignies Hub — nouvelle demande {labels.get(x.type,x.type)}"
        msg["From"]=sender
        msg["To"]=", ".join(recipients)
        msg.set_content(f"""Nouvelle demande envoyée par {member.first_name} {member.last_name}.

Type : {labels.get(x.type,x.type)}
Sujet : {x.title or x.guest_name or "—"}
Message :
{x.message or "—"}

Cette demande est également disponible dans BNI Soignies Hub.""")
        if x.attachment_data:
            maintype,subtype=(x.attachment_type or "application/octet-stream").split("/",1) if "/" in (x.attachment_type or "") else ("application","octet-stream")
            msg.add_attachment(x.attachment_data,maintype=maintype,subtype=subtype,filename=x.attachment_name or "piece-jointe")
        port=int(os.getenv("SMTP_PORT","587"))
        user=os.getenv("SMTP_USER","").strip();password=os.getenv("SMTP_PASSWORD","")
        use_ssl=os.getenv("SMTP_SSL","false").lower() in {"1","true","yes"}
        smtp=(smtplib.SMTP_SSL(host,port,timeout=15) if use_ssl else smtplib.SMTP(host,port,timeout=15))
        try:
            if not use_ssl and os.getenv("SMTP_STARTTLS","true").lower() in {"1","true","yes"}:smtp.starttls()
            if user:smtp.login(user,password)
            smtp.send_message(msg)
            return True
        finally:smtp.quit()

    @app.route("/api/me/requests",methods=["GET","POST"])
    def my_requests():
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        if request.method=="GET":
            rows=MemberRequest.query.filter_by(member_id=account.member_id).order_by(MemberRequest.created_at.desc()).all()
            return jsonify([member_request_payload(x) for x in rows])
        data=request.form if request.content_type and request.content_type.startswith("multipart/form-data") else (request.get_json(silent=True) or {})
        kind=(data.get("type") or "").strip().upper()
        allowed={"INFOMERCIAL","COMMUNICATION","TRAINING","MENTORING","EVENT","INVITATION"}
        if kind not in allowed:return {"error":"Type de demande invalide."},400
        event_date=None
        if data.get("eventDate"):
            try:event_date=date.fromisoformat(data["eventDate"])
            except ValueError:return {"error":"Date invalide."},400
        upload=request.files.get("attachment")
        attachment_name=attachment_type=None;attachment_data=None
        if upload and upload.filename:
            attachment_name=upload.filename
            ext=attachment_name.rsplit(".",1)[-1].lower() if "." in attachment_name else ""
            if ext not in REQUEST_ALLOWED_EXTENSIONS:return {"error":"Type de pièce jointe non autorisé. Utilisez une image, un PDF, un PowerPoint ou une mini-vidéo."},400
            attachment_data=upload.read(REQUEST_MAX_FILE_SIZE+1)
            if len(attachment_data)>REQUEST_MAX_FILE_SIZE:return {"error":"La pièce jointe dépasse 5 Mo."},413
            attachment_type=upload.mimetype or "application/octet-stream"
        x=MemberRequest(member_id=account.member_id,type=kind,title=(data.get("title") or "").strip() or None,message=(data.get("message") or "").strip() or None,guest_name=(data.get("guestName") or "").strip() or None,guest_email=(data.get("guestEmail") or "").strip() or None,event_date=event_date,attachment_name=attachment_name,attachment_type=attachment_type,attachment_data=attachment_data,status="TO_PROCESS")
        db.session.add(x);db.session.flush()
        targets=request_assignees(kind)
        for target in targets:
            if target["memberId"]!=account.member_id:notify(target["memberId"],"NEW_REQUEST",f"Nouvelle demande : {kind.title()}",f"{account.member.first_name} {account.member.last_name} a envoyé une nouvelle demande.","requests")
        audit(account,"MEMBER_REQUEST_CREATED","member_request",x.id,f"type={kind}; attachment={attachment_name or 'none'}")
        db.session.commit()
        email_sent=False
        try:email_sent=send_request_email(x,account,targets)
        except Exception as exc:
            app.logger.exception("Request email delivery failed: %s",exc)
        payload=member_request_payload(x);payload["emailSent"]=email_sent
        return payload,201

    @app.get("/api/member-requests/<int:request_id>/attachment")
    def member_request_attachment(request_id):
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        x=db.session.get(MemberRequest,request_id)
        if not x or not x.attachment_data:return {"error":"Pièce jointe introuvable."},404
        if not (account.is_system_admin or account.member_id==x.member_id or can_manage_request(account,x.type)):return {"error":"Accès refusé."},403
        return send_file(io.BytesIO(x.attachment_data),mimetype=x.attachment_type or "application/octet-stream",download_name=x.attachment_name or "piece-jointe",as_attachment=True)

    @app.route("/api/request-recipients",methods=["GET","PUT"])
    def request_recipients_config():
        account,err=require_admin()
        if err:return err
        allowed={"INFOMERCIAL","COMMUNICATION","TRAINING","MENTORING","EVENT","INVITATION"}
        if request.method=="GET":
            out={}
            for kind in allowed:
                configured=RequestRecipient.query.filter_by(request_type=kind).order_by(RequestRecipient.id).all()
                out[kind]=[x.member_id for x in configured]
            return jsonify(out)
        data=request.get_json(silent=True) or {}
        kind=(data.get("type") or "").strip().upper()
        if kind not in allowed:return {"error":"Type de demande invalide."},400
        raw=data.get("memberIds") or []
        try:mids=list(dict.fromkeys([int(x) for x in raw]))
        except (TypeError,ValueError):return {"error":"Liste de membres invalide."},400
        if mids:
            valid={m.id for m in Member.query.filter(Member.id.in_(mids)).all()}
            if len(valid)!=len(mids):return {"error":"Un ou plusieurs membres sont introuvables."},400
        RequestRecipient.query.filter_by(request_type=kind).delete()
        for mid in mids:db.session.add(RequestRecipient(request_type=kind,member_id=mid))
        audit(account,"REQUEST_RECIPIENTS_CHANGED","request_routing",kind,f"memberIds={mids}")
        db.session.commit()
        return {"type":kind,"memberIds":mids,"assignees":[{k:v for k,v in x.items() if k!="email"} for x in request_assignees(kind)]}

    @app.get("/api/member-requests")
    def all_member_requests():
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        rows=MemberRequest.query.order_by(MemberRequest.created_at.desc()).all()
        if not account.is_system_admin:
            rows=[x for x in rows if can_manage_request(account,x.type)]
        return jsonify([member_request_payload(x) for x in rows])

    @app.patch("/api/member-requests/<int:request_id>")
    def update_member_request(request_id):
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        x=db.session.get(MemberRequest,request_id)
        if not x:return {"error":"Demande introuvable."},404
        if not can_manage_request(account,x.type):return {"error":"Cette demande relève d'une autre fonction."},403
        status=(request.get_json(silent=True) or {}).get("status")
        if status not in {"TO_PROCESS","IN_PROGRESS","COMPLETED"}:return {"error":"Statut invalide."},400
        x.status=status
        audit(account,"MEMBER_REQUEST_STATUS_CHANGED","member_request",x.id,f"status={status}")
        if x.member_id!=account.member_id:notify(x.member_id,"REQUEST_STATUS","Mise à jour de votre demande",f"Votre demande est maintenant : {status.replace('_',' ').lower()}.","my-requests")
        db.session.commit();return member_request_payload(x)

    @app.get("/api/me/requests/summary")
    def my_requests_summary():
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        rows=MemberRequest.query.filter_by(member_id=account.member_id).order_by(MemberRequest.created_at.desc()).all()
        counts={"TO_PROCESS":0,"IN_PROGRESS":0,"COMPLETED":0}
        for x in rows:counts[x.status]=counts.get(x.status,0)+1
        return {"total":len(rows),"counts":counts,"recent":[member_request_payload(x) for x in rows[:5]]}

    @app.get("/api/me/request-inbox")
    def my_request_inbox():
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        rows=MemberRequest.query.order_by(MemberRequest.created_at.desc()).all()
        return jsonify([member_request_payload(x) for x in rows if can_manage_request(account,x.type)])

    @app.get("/api/me/duty-assignments")
    def my_duty_assignments():
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        rows=DutyAssignment.query.join(Meeting,DutyAssignment.meeting_id==Meeting.id).filter(DutyAssignment.member_id==account.member_id,DutyAssignment.status.in_(["PUBLISHED","TO_REASSIGN"])).order_by(Meeting.date).all()
        return jsonify([assignment_payload(a) for a in rows])

    @app.post("/api/me/duty-assignments/<int:assignment_id>/unavailable")
    def my_duty_unavailable(assignment_id):
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        a=db.session.get(DutyAssignment,assignment_id)
        if not a or a.member_id!=account.member_id:return {"error":"Affectation introuvable."},404
        if a.status!="PUBLISHED":return {"error":"Cette permanence n'est pas publiée."},400
        a.status="TO_REASSIGN";existing=SwapRequest.query.filter_by(assignment_id=a.id,status="OPEN").first()
        if not existing:db.session.add(SwapRequest(assignment_id=a.id,requester_id=account.member_id,status="OPEN"))
        meeting=db.session.get(Meeting,a.meeting_id);duty=db.session.get(Duty,a.duty_id)
        for m in Member.query.all():
            ms=Membership.query.filter_by(member_id=m.id).order_by(Membership.id.desc()).first()
            if m.id!=account.member_id and ms and ms.status=="ACTIVE" and (not ms.start_date or ms.start_date<=meeting.date) and (not ms.end_date or ms.end_date>=meeting.date):
                notify(m.id,"DUTY_REPLACEMENT","Remplacement recherché",f"{account.member.first_name} {account.member.last_name} cherche un remplacement pour {duty.name} le {meeting.date.strftime('%d/%m/%Y')}.","duties")
        db.session.commit()
        return {"status":"TO_REASSIGN","message":"Indisponibilité enregistrée. Les autres membres peuvent proposer de reprendre cette permanence."}

    @app.get("/api/me/profile")
    def my_profile():
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        m=account.member
        return {"id":m.id,"firstName":m.first_name,"lastName":m.last_name,"company":m.company,"activity":m.activity,"email":account.email,"phone":m.phone,"roles":account_payload(account)["roles"]}

    def swap_payload(x):
        a=db.session.get(DutyAssignment,x.assignment_id);meeting=db.session.get(Meeting,a.meeting_id) if a else None;duty=db.session.get(Duty,a.duty_id) if a else None;requester=db.session.get(Member,x.requester_id);volunteer=db.session.get(Member,x.volunteer_id) if x.volunteer_id else None
        return {"id":x.id,"assignmentId":x.assignment_id,"date":meeting.date.isoformat() if meeting else None,"duty":duty.name if duty else "—","requester":f"{requester.first_name} {requester.last_name}" if requester else "—","volunteer":f"{volunteer.first_name} {volunteer.last_name}" if volunteer else None,"status":x.status}

    @app.get("/api/swaps/open")
    def open_swaps():
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        rows=SwapRequest.query.filter_by(status="OPEN").order_by(SwapRequest.created_at).all()
        return jsonify([swap_payload(x) for x in rows if x.requester_id!=account.member_id])

    @app.post("/api/swaps/<int:swap_id>/volunteer")
    def volunteer_swap(swap_id):
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        x=db.session.get(SwapRequest,swap_id)
        if not x or x.status!="OPEN":return {"error":"Cette demande n'est plus disponible."},404
        if x.requester_id==account.member_id:return {"error":"Vous ne pouvez pas reprendre votre propre permanence."},400
        a=db.session.get(DutyAssignment,x.assignment_id)
        if DutyAssignment.query.filter(DutyAssignment.meeting_id==a.meeting_id,DutyAssignment.member_id==account.member_id,DutyAssignment.id!=a.id).first():return {"error":"Vous avez déjà une fonction lors de cette réunion."},409
        meeting=db.session.get(Meeting,a.meeting_id)
        ms=Membership.query.filter_by(member_id=account.member_id).order_by(Membership.id.desc()).first()
        if not ms or ms.status!="ACTIVE" or (ms.start_date and ms.start_date>meeting.date) or (ms.end_date and ms.end_date<meeting.date):return {"error":"Vous n'êtes pas membre actif à cette date."},403
        requester=x.requester_id;x.volunteer_id=account.member_id;x.status="ACCEPTED";a.member_id=account.member_id;a.status="PUBLISHED"
        notify(requester,"DUTY_REPLACED","Remplacement trouvé",f"{account.member.first_name} {account.member.last_name} reprend votre permanence du {meeting.date.strftime('%d/%m/%Y')}.","duties")
        notify(account.member_id,"DUTY_ASSIGNED","Permanence reprise",f"Vous êtes maintenant titulaire de cette permanence du {meeting.date.strftime('%d/%m/%Y')}.","duties")
        db.session.commit()
        return swap_payload(x)

    @app.patch("/api/duty-assignments/<int:assignment_id>")
    def update_duty_assignment(assignment_id):
        _,auth_error=require_admin()
        if auth_error:return auth_error
        a=db.session.get(DutyAssignment,assignment_id)
        if not a:return {"error":"Affectation introuvable."},404
        data=request.get_json(silent=True) or {}
        if "memberId" in data:
            try:member_id=int(data["memberId"])
            except (TypeError,ValueError):return {"error":"Membre invalide."},400
            member=db.session.get(Member,member_id)
            if not member:return {"error":"Membre introuvable."},404
            meeting=db.session.get(Meeting,a.meeting_id)
            ms=Membership.query.filter_by(member_id=member_id).order_by(Membership.id.desc()).first()
            if not ms or ms.status!="ACTIVE" or (ms.start_date and ms.start_date>meeting.date) or (ms.end_date and ms.end_date<meeting.date):return {"error":"Membre non actif à cette date."},400
            if DutyAssignment.query.filter(DutyAssignment.meeting_id==a.meeting_id,DutyAssignment.member_id==member_id,DutyAssignment.id!=a.id).first():return {"error":"Ce membre a déjà une fonction ce jour-là."},409
            previous=a.member_id;a.member_id=member_id
            SwapRequest.query.filter_by(assignment_id=a.id,status="OPEN").update({"status":"CANCELLED"})
            if previous!=member_id:
                notify(previous,"DUTY_CHANGED","Permanence réattribuée",f"Votre permanence du {meeting.date.strftime('%d/%m/%Y')} a été réattribuée.","duties")
                notify(member_id,"DUTY_ASSIGNED","Nouvelle permanence",f"Une permanence vous a été attribuée le {meeting.date.strftime('%d/%m/%Y')}.","duties")
        db.session.commit();return assignment_payload(a)

    @app.post("/api/duty-assignments/<int:assignment_id>/reassign")
    def reassign_duty_assignment(assignment_id):
        _,auth_error=require_admin()
        if auth_error:return auth_error
        a=db.session.get(DutyAssignment,assignment_id)
        if not a:return {"error":"Affectation introuvable."},404
        meeting=db.session.get(Meeting,a.meeting_id);duty=db.session.get(Duty,a.duty_id)
        occupied={x.member_id for x in DutyAssignment.query.filter_by(meeting_id=meeting.id).all() if x.id!=a.id}
        candidates=[]
        for m in Member.query.order_by(Member.last_name,Member.first_name).all():
            ms=Membership.query.filter_by(member_id=m.id).order_by(Membership.id.desc()).first()
            if ms and ms.status=="ACTIVE" and m.id not in occupied and (not ms.start_date or ms.start_date<=meeting.date) and (not ms.end_date or ms.end_date>=meeting.date):candidates.append(m)
        if not candidates:return {"error":"Aucun membre disponible."},400
        def score(m):
            total=DutyAssignment.query.filter_by(member_id=m.id).count()
            specific=DutyAssignment.query.filter_by(member_id=m.id,duty_id=duty.id).count()
            return (total,specific,m.last_name.lower(),m.first_name.lower())
        candidates=[m for m in candidates if m.id!=a.member_id] or candidates
        candidates.sort(key=score);previous=a.member_id;a.member_id=candidates[0].id
        SwapRequest.query.filter_by(assignment_id=a.id,status="OPEN").update({"status":"CANCELLED"})
        notify(previous,"DUTY_CHANGED","Permanence réattribuée",f"Votre permanence du {meeting.date.strftime('%d/%m/%Y')} a été réattribuée.","duties")
        notify(a.member_id,"DUTY_ASSIGNED","Nouvelle permanence",f"Une permanence vous a été attribuée le {meeting.date.strftime('%d/%m/%Y')}.","duties")
        audit(_, "DUTY_REASSIGNED","duty_assignment",a.id,f"from={previous}; to={a.member_id}")
        db.session.commit();return assignment_payload(a)

    @app.get("/api/planning/welcome")
    def welcome_planning():
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        rows=DutyAssignment.query.join(Meeting,DutyAssignment.meeting_id==Meeting.id).filter(DutyAssignment.status.in_(["PUBLISHED","TO_REASSIGN"])).order_by(Meeting.date,DutyAssignment.duty_id).all()
        grouped={}
        for a in rows:
            p=assignment_payload(a);key=p["date"]
            grouped.setdefault(key,{"date":key,"assignments":[]})["assignments"].append(p)
        return jsonify(list(grouped.values()))

    @app.post("/api/duty-assignments/generate")
    def generate_duty_assignments():
        _,auth_error=require_admin()
        if auth_error:return auth_error
        data=request.get_json(silent=True) or {}
        try:start=date.fromisoformat(data.get("startDate",""));end=date.fromisoformat(data.get("endDate",""))
        except ValueError:return {"error":"Dates invalides."},400
        duties=Duty.query.filter_by(active=True).order_by(Duty.id).all()
        if not duties:return {"error":"Aucune fonction d'accueil configurée."},400
        meetings=Meeting.query.filter(Meeting.date>=start,Meeting.date<=end,Meeting.type!="NO_MEETING").order_by(Meeting.date).all()
        active_members=[]
        for m in Member.query.order_by(Member.last_name,Member.first_name).all():
            ms=Membership.query.filter_by(member_id=m.id).order_by(Membership.id.desc()).first()
            if ms and ms.status=="ACTIVE" and (not ms.end_date or ms.end_date>=start):active_members.append(m)
        if len(active_members)<len(duties):return {"error":"Pas assez de membres actifs pour générer la rotation."},400
        existing=DutyAssignment.query.join(Meeting,DutyAssignment.meeting_id==Meeting.id).filter(Meeting.date>=start,Meeting.date<=end).all()
        total={m.id:0 for m in active_members};per={(m.id,d.id):0 for m in active_members for d in duties};last={m.id:date.min for m in active_members}
        for a in existing:
            if a.member_id in total:
                total[a.member_id]+=1;per[(a.member_id,a.duty_id)]=per.get((a.member_id,a.duty_id),0)+1
                mt=db.session.get(Meeting,a.meeting_id)
                if mt and mt.date>last[a.member_id]:last[a.member_id]=mt.date
        created=0
        for mt in meetings:
            used={a.member_id for a in existing if a.meeting_id==mt.id}
            for duty in duties:
                if any(a.meeting_id==mt.id and a.duty_id==duty.id for a in existing):continue
                candidates=[m for m in active_members if m.id not in used and (not Membership.query.filter_by(member_id=m.id).order_by(Membership.id.desc()).first().start_date or Membership.query.filter_by(member_id=m.id).order_by(Membership.id.desc()).first().start_date<=mt.date)]
                if not candidates:continue
                candidates.sort(key=lambda m:(total[m.id],per.get((m.id,duty.id),0),last[m.id],m.last_name.lower(),m.first_name.lower()))
                chosen=candidates[0]
                a=DutyAssignment(meeting_id=mt.id,duty_id=duty.id,member_id=chosen.id,status="DRAFT");db.session.add(a);existing.append(a);used.add(chosen.id);total[chosen.id]+=1;per[(chosen.id,duty.id)]=per.get((chosen.id,duty.id),0)+1;last[chosen.id]=mt.date;created+=1
        db.session.commit()
        return {"status":"ok","created":created}

    @app.post("/api/duty-assignments/publish")
    def publish_duty_assignments():
        _,auth_error=require_admin()
        if auth_error:return auth_error
        data=request.get_json(silent=True) or {}
        try:start=date.fromisoformat(data.get("startDate",""));end=date.fromisoformat(data.get("endDate",""))
        except ValueError:return {"error":"Dates invalides."},400
        rows=DutyAssignment.query.join(Meeting,DutyAssignment.meeting_id==Meeting.id).filter(Meeting.date>=start,Meeting.date<=end,DutyAssignment.status=="DRAFT").all()
        for a in rows:a.status="PUBLISHED"
        audit(_, "DUTY_ASSIGNMENTS_PUBLISHED","duty_assignment",None,f"{start.isoformat()} to {end.isoformat()}; count={len(rows)}")
        db.session.commit();return {"status":"ok","published":len(rows)}

    @app.route("/api/meetings",methods=["GET","POST"])
    def meetings():
        account,_=bearer_account()
        if not account:return {"error":"Authentification requise."},401
        if request.method=="POST":
            _,auth_error=require_admin()
            if auth_error:return auth_error
            data=request.get_json(silent=True) or {}
            try: meeting_date=date.fromisoformat(data.get("date",""))
            except ValueError:return {"error":"Date invalide."},400
            allowed={"REGULAR","NO_MEETING","EXTERNAL","HOSTED","JOINT","SPECIAL"}
            meeting_type=data.get("type","REGULAR")
            if meeting_type not in allowed:return {"error":"Type de réunion invalide."},400
            if Meeting.query.filter_by(date=meeting_date).first():return {"error":"Une entrée existe déjà à cette date."},409
            start=None
            if data.get("startTime"):
                try:start=datetime.strptime(data["startTime"],"%H:%M").time()
                except ValueError:return {"error":"Heure invalide."},400
            m=Meeting(date=meeting_date,start_time=start,type=meeting_type,location=(data.get("location") or "").strip() or None,status=data.get("status","PLANNED"),notes=(data.get("notes") or "").strip() or None)
            db.session.add(m);db.session.flush();audit(_, "MEETING_CREATED","meeting",m.id,f"date={meeting_date.isoformat()}; type={meeting_type}");db.session.commit()
            return meeting_payload(m),201
        rows=Meeting.query.order_by(Meeting.date).all()
        return jsonify([meeting_payload(m) for m in rows])

    @app.route("/api/meetings/<int:meeting_id>",methods=["PATCH","DELETE"])
    def meeting_detail(meeting_id):
        _,auth_error=require_admin()
        if auth_error:return auth_error
        m=db.session.get(Meeting,meeting_id)
        if not m:return {"error":"Réunion introuvable."},404
        if request.method=="DELETE":
            audit(_, "MEETING_DELETED","meeting",m.id,f"date={m.date.isoformat()}");db.session.delete(m);db.session.commit();return {"status":"ok"}
        data=request.get_json(silent=True) or {}
        if "date" in data:
            try:m.date=date.fromisoformat(data["date"])
            except ValueError:return {"error":"Date invalide."},400
        if "type" in data:m.type=data["type"]
        if "location" in data:m.location=(data["location"] or "").strip() or None
        if "notes" in data:m.notes=(data["notes"] or "").strip() or None
        if "status" in data:m.status=data["status"]
        if "startTime" in data:
            try:m.start_time=datetime.strptime(data["startTime"],"%H:%M").time() if data["startTime"] else None
            except ValueError:return {"error":"Heure invalide."},400
        audit(_, "MEETING_UPDATED","meeting",m.id,f"date={m.date.isoformat()}; type={m.type}");db.session.commit();return meeting_payload(m)

    @app.post("/api/meetings/generate-thursdays")
    def generate_thursdays():
        _,auth_error=require_admin()
        if auth_error:return auth_error
        data=request.get_json(silent=True) or {}
        try:start=date.fromisoformat(data.get("startDate",""));end=date.fromisoformat(data.get("endDate",""))
        except ValueError:return {"error":"Dates invalides."},400
        if end<start:return {"error":"La date de fin doit suivre la date de début."},400
        d=start
        while d.weekday()!=3:d+=timedelta(days=1)
        created=0
        while d<=end:
            if not Meeting.query.filter_by(date=d).first():
                db.session.add(Meeting(date=d,start_time=datetime.strptime(data.get("startTime","07:00"),"%H:%M").time(),type="REGULAR",location=(data.get("location") or "").strip() or None,status="PLANNED"));created+=1
            d+=timedelta(days=7)
        db.session.commit()
        return {"status":"ok","created":created}

    with app.app_context():
        db.create_all()
        # Temporary lightweight migration until Alembic is introduced.
        db.session.execute(text("ALTER TABLE user_accounts ADD COLUMN IF NOT EXISTS is_system_admin BOOLEAN NOT NULL DEFAULT FALSE"))
        db.session.execute(text("ALTER TABLE member_requests ADD COLUMN IF NOT EXISTS attachment_name VARCHAR(255)"))
        db.session.execute(text("ALTER TABLE member_requests ADD COLUMN IF NOT EXISTS attachment_type VARCHAR(120)"))
        db.session.execute(text("ALTER TABLE member_requests ADD COLUMN IF NOT EXISTS attachment_data BYTEA"))
        db.session.commit()
        for code,name in [("SETUP","Préparation"),("RECEPTION","Accueil des invités"),("FOLLOWUP","Suivi des invités")]:
            duty=Duty.query.filter_by(code=code).first()
            if not duty:db.session.add(Duty(code=code,name=name,active=True))
        db.session.commit()
        # Idempotent bootstrap of the first system administrator.
        admin_email=os.getenv("SYSTEM_ADMIN_EMAIL","").strip().lower()
        if admin_email:
            account=UserAccount.query.filter(db.func.lower(UserAccount.email)==admin_email).first()
            if account:
                account.is_system_admin=True
                membership=Membership.query.filter_by(member_id=account.member_id).order_by(Membership.id.asc()).first()
                if not membership:
                    membership=Membership(member_id=account.member_id,status="ACTIVE")
                    db.session.add(membership)
                membership.start_date=date(2017,4,1)
                membership.end_date=None
                membership.status="ACTIVE"
                mandate=Mandate.query.filter_by(name="Octobre 2026 - Mars 2027").first()
                if not mandate:
                    mandate=Mandate(name="Octobre 2026 - Mars 2027",start_date=date(2026,10,1),end_date=date(2027,3,31),status="PREPARATION")
                    db.session.add(mandate);db.session.flush()
                role=Role.query.filter(db.func.lower(Role.name)=="président").first()
                if not role:
                    role=Role(name="Président");db.session.add(role);db.session.flush()
                assignment=RoleAssignment.query.filter_by(mandate_id=mandate.id,role_id=role.id,member_id=account.member_id).first()
                if not assignment:
                    db.session.add(RoleAssignment(mandate_id=mandate.id,role_id=role.id,member_id=account.member_id,status="ACTIVE"))
                db.session.commit()
    return app

app=create_app()

if __name__=="__main__":
    app.run(host="0.0.0.0",port=int(os.getenv("PORT","5000")))
